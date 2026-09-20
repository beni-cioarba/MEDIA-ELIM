"""
Importa el plan de lectura bíblica desde el Excel de la iglesia y genera
`src/app/core/bible-reading.config.ts`.

Uso:
    python scripts/import-bible-plan.py "ruta/PROGRAMARE_CITIREA_BIBLIEI.xlsx"

Estructura esperada del Excel (una hoja por mes, p. ej. «2026 OCTOMBRIE»):
    A1: año        B1: título del plan («PROGRAMARE CITIREA BIBLIEI - VDC»)
    A2: MES        B2: tramo del mes («Ieremia 19-52, Plângerile…»)
    A3: cabeceras  Săptămâna | Data | Vechiul testament | Noul testament
    A4…: por semana, 7 filas (lunes→domingo):
         · en la primera fila, A = número de semana
         · en la tercera fila, A = resumen del tramo semanal (celda combinada)
         · B = fecha, C = lectura AT, D = lectura NT (opcional)

El resumen semanal se toma de la celda combinada si es texto; si falta o no
es texto (hay hojas con un número suelto ahí), se deriva de las lecturas
diarias para que la proyección nunca se quede sin tramo.

Sólo necesita `openpyxl` (viene con la instalación de Python usada en el
proyecto para el resto de scripts de datos).
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl

OUTPUT = Path(__file__).resolve().parents[1] / "src" / "app" / "core" / "bible-reading.config.ts"

MONTHS_RO = {
    "IANUARIE": 1, "FEBRUARIE": 2, "MARTIE": 3, "APRILIE": 4, "MAI": 5, "IUNIE": 6,
    "IULIE": 7, "AUGUST": 8, "SEPTEMBRIE": 9, "OCTOMBRIE": 10, "NOIEMBRIE": 11, "DECEMBRIE": 12,
}


@dataclass
class Week:
    number: int
    start: date
    summary: str | None
    readings: list[str] = field(default_factory=list)
    nt_readings: list[str] = field(default_factory=list)


@dataclass
class Month:
    id: str
    label: str
    scope: str
    weeks: list[Week]


def as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def ts_string(value: str) -> str:
    return "'" + value.replace("\\", "\\\\").replace("'", "\\'") + "'"


def derive_summary(readings: list[str]) -> str:
    """«Iona 1-2», «Iona 3-4», «Mica 1-2» → «Iona 1-4, Mica 1-2»."""
    ranges: list[tuple[str, int, int]] = []
    for reading in readings:
        for part in reading.split(","):
            match = re.match(r"^\s*(.+?)\s+(\d+)(?:-(\d+))?\s*$", part)
            if not match:
                continue
            book, first, last = match.group(1), int(match.group(2)), int(match.group(3) or match.group(2))
            if ranges and ranges[-1][0] == book:
                ranges[-1] = (book, ranges[-1][1], max(ranges[-1][2], last))
            else:
                ranges.append((book, first, last))
    return ", ".join(f"{book} {a}" if a == b else f"{book} {a}-{b}" for book, a, b in ranges)


def read_month(ws) -> Month:
    year = int(ws["A1"].value)
    month_name = str(ws["A2"].value).strip().upper()
    month_number = MONTHS_RO[month_name]
    scope = str(ws["B2"].value or "").strip()

    weeks: list[Week] = []
    current: Week | None = None
    for row in ws.iter_rows(min_row=4, values_only=True):
        col_a, col_b, col_c, col_d = (list(row) + [None] * 4)[:4]
        day = as_date(col_b)
        if day is None:
            continue
        if isinstance(col_a, (int, float)) and current is not None and int(col_a) != current.number and len(current.readings) >= 6:
            current = None
        if current is None:
            if not isinstance(col_a, (int, float)):
                raise ValueError(f"{ws.title}: la fila del {day} debería abrir una semana con su número en A")
            current = Week(number=int(col_a), start=day, summary=None)
            weeks.append(current)
        elif isinstance(col_a, str) and col_a.strip():
            current.summary = col_a.strip()
        current.readings.append(str(col_c or "").strip())
        if col_d:
            current.nt_readings.append(str(col_d).strip())

    for week in weeks:
        if not week.summary:
            week.summary = derive_summary(week.readings)

    return Month(
        id=f"{year}-{month_number:02d}",
        label=f"{month_name.capitalize()} {year}",
        scope=scope,
        weeks=weeks,
    )


def render(months: list[Month], title: str, source: str) -> str:
    lines: list[str] = []
    lines.append("/**")
    lines.append(" * Plan de lectura bíblica — GENERADO desde el Excel de la iglesia.")
    lines.append(" *")
    lines.append(" * No lo edites a mano: ejecuta")
    lines.append(' *   python scripts/import-bible-plan.py "ruta/PROGRAMARE_CITIREA_BIBLIEI.xlsx"')
    lines.append(" * y el fichero se regenera entero (un mes = una hoja del Excel).")
    lines.append(" *")
    lines.append(f" * Fuente: {source}")
    lines.append(" */")
    lines.append("")
    lines.append("/** Semana del plan: empieza en lunes; `readings[i]` es la lectura del día `start + i`. */")
    lines.append("export interface BibleReadingWeek {")
    lines.append("  /** Número de semana dentro del plan plurianual (90, 91…). */")
    lines.append("  readonly number: number;")
    lines.append("  /** Lunes de la semana, `YYYY-MM-DD`. */")
    lines.append("  readonly start: string;")
    lines.append("  /** Tramo de la semana, tal como se anuncia («Isaia 57-66, Ieremia 1-4»). */")
    lines.append("  readonly summary: string;")
    lines.append("  /** Lecturas diarias, lunes→domingo (puede haber menos de 7 al cerrar el plan). */")
    lines.append("  readonly readings: readonly string[];")
    lines.append("  /** Lecturas del Nuevo Testamento del mismo día, si el plan las trae. */")
    lines.append("  readonly newTestament?: readonly string[];")
    lines.append("}")
    lines.append("")
    lines.append("/** Un mes del plan = una hoja del Excel. */")
    lines.append("export interface BibleReadingMonth {")
    lines.append("  /** `YYYY-MM`. */")
    lines.append("  readonly id: string;")
    lines.append("  /** Rótulo en rumano («Octombrie 2026»). */")
    lines.append("  readonly label: string;")
    lines.append("  /** Tramo del mes completo. */")
    lines.append("  readonly scope: string;")
    lines.append("  readonly weeks: readonly BibleReadingWeek[];")
    lines.append("}")
    lines.append("")
    lines.append("export interface BibleReadingPlan {")
    lines.append("  /** Título del plan («Programare citirea Bibliei – VDC»). */")
    lines.append("  readonly title: string;")
    lines.append("  readonly months: readonly BibleReadingMonth[];")
    lines.append("}")
    lines.append("")
    lines.append("export const BIBLE_READING_PLAN: BibleReadingPlan = {")
    lines.append(f"  title: {ts_string(title)},")
    lines.append("  months: [")
    for month in months:
        lines.append("    {")
        lines.append(f"      id: {ts_string(month.id)},")
        lines.append(f"      label: {ts_string(month.label)},")
        lines.append(f"      scope: {ts_string(month.scope)},")
        lines.append("      weeks: [")
        for week in month.weeks:
            lines.append("        {")
            lines.append(f"          number: {week.number},")
            lines.append(f"          start: {ts_string(week.start.isoformat())},")
            lines.append(f"          summary: {ts_string(week.summary or '')},")
            lines.append("          readings: [")
            for reading in week.readings:
                lines.append(f"            {ts_string(reading)},")
            lines.append("          ],")
            if week.nt_readings:
                lines.append("          newTestament: [")
                for reading in week.nt_readings:
                    lines.append(f"            {ts_string(reading)},")
                lines.append("          ],")
            lines.append("        },")
        lines.append("      ],")
        lines.append("    },")
    lines.append("  ],")
    lines.append("};")
    lines.append("")
    return "\n".join(lines)


def main() -> None:
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    source = Path(sys.argv[1])
    wb = openpyxl.load_workbook(source, data_only=True)

    months = [read_month(ws) for ws in wb.worksheets]
    months.sort(key=lambda m: m.id)

    # «PROGRAMARE CITIREA BIBLIEI - VDC» → «Programare citirea Bibliei – VDC».
    raw_title = str(wb.worksheets[0]["B1"].value or "Programare citirea Bibliei").strip()
    title = re.sub(r"\s*-\s*", " – ", raw_title.capitalize())
    title = re.sub(r"\bbibliei\b", "Bibliei", title)
    title = re.sub(r"\bvdc\b", "VDC", title)

    # Saltos LF explícitos: el repositorio los normaliza y así el diff es limpio.
    with OUTPUT.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write(render(months, title, source.name))
    total_weeks = sum(len(m.weeks) for m in months)
    print(f"OK: {OUTPUT.relative_to(OUTPUT.parents[3])} — {len(months)} meses, {total_weeks} semanas")
    for m in months:
        numbers = [w.number for w in m.weeks]
        print(f"  {m.label}: semanas {numbers[0]}–{numbers[-1]} · {m.scope}")


if __name__ == "__main__":
    main()
